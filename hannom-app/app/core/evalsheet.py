"""The team's "format đánh giá" spreadsheet.

Eight columns in the order they already use, and the point of the format: every
character where **Label** (their transcription) and **Corrected** (what the
reviewer read) disagree is coloured — red on the Label side, blue on Corrected —
so a mistake is visible without comparing two blocks of CJK by eye.

Both accuracy columns divide by ``max(len(a), len(b))``, which is their
Task.xlsx convention rather than the CER used elsewhere in this app. Checked
against their worked example: 78.57% = 1 - 6/28 with whitespace removed, and
80.65% = 1 - 6/31 with the three line breaks counted.

Note the deliberate gap with ``reviews.xlsx``, which reports CER. That file is
for measuring; this one is for reading alongside their existing sheets, so it
matches their arithmetic.
"""

from __future__ import annotations

import re
import unicodedata
from difflib import SequenceMatcher
from typing import Any

from app.core.metrics import levenshtein

HEADERS = [
    "Post ID",
    "Image",
    "FB Caption",
    "Label",
    "Corrected",
    "Levenshtein Accuracy",
    "Levenshtein Accuracy (bao gồm cả line break)",
    "Note",
]
WIDTHS = [46, 46, 52, 30, 30, 15, 20, 44]

# Verdicts that produce no corrected text. Scoring them 0% would drag the
# averages down with images nobody could judge, so they are left blank and the
# reason moves into the Note — this format has no verdict column of its own.
NO_TRANSCRIPTION = {
    "unreadable": "Không đọc được",
    "not_an_image": "Không phải ảnh thư pháp",
}

_WS = re.compile(r"\s+")


def fold(text: str, keep_breaks: bool) -> str:
    """NFC, then either drop all whitespace or keep the line structure."""
    folded = unicodedata.normalize("NFC", text or "")
    if keep_breaks:
        return folded.replace("\r\n", "\n").replace("\r", "\n").strip()
    return _WS.sub("", folded)


def accuracy(label: str, corrected: str, *, keep_breaks: bool) -> float | None:
    """Their convention: 1 - distance / max(len). Not CER."""
    a, b = fold(label, keep_breaks), fold(corrected, keep_breaks)
    if not a and not b:
        return 1.0
    longest = max(len(a), len(b))
    if not longest:
        return 1.0
    return max(0.0, 1.0 - levenshtein(a, b) / longest)


def diff_spans(label: str, corrected: str) -> tuple[list[tuple[str, bool]], list[tuple[str, bool]]]:
    """Split both strings into (text, differs) runs.

    Aligned with SequenceMatcher rather than compared position by position: a
    single inserted character would otherwise paint the whole rest of the line
    as wrong, which is worse than no highlighting at all.
    """
    left: list[tuple[str, bool]] = []
    right: list[tuple[str, bool]] = []

    for tag, i1, i2, j1, j2 in SequenceMatcher(
        None, label, corrected, autojunk=False
    ).get_opcodes():
        same = tag == "equal"
        if label[i1:i2]:
            left.append((label[i1:i2], not same))
        if corrected[j1:j2]:
            right.append((corrected[j1:j2], not same))

    return left, right


def _rich(spans: list[tuple[str, bool]], font) -> Any:
    """A rich-text cell value, or a plain string when nothing differs.

    openpyxl writes an empty CellRichText as a blank cell, so a string is the
    safer representation when there is no highlighting to apply.
    """
    from openpyxl.cell.rich_text import CellRichText, TextBlock

    if not any(differs for _, differs in spans):
        return "".join(text for text, _ in spans)

    value = CellRichText()
    for text, differs in spans:
        value.append(TextBlock(font, text) if differs else text)
    return value


def build(rows: list[dict[str, Any]]) -> Any:
    """An ``openpyxl`` Workbook of the reviewed rows, in their format."""
    from openpyxl import Workbook
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    red = InlineFont(color="FFC00000", b=True)
    blue = InlineFont(color="FF0070C0", b=True)

    book = Workbook()
    sheet = book.active
    sheet.title = "Đánh giá"

    header_fill = PatternFill("solid", fgColor="FFBFBFBF")
    thin = Side(style="thin", color="FF808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top_wrap = Alignment(vertical="top", wrap_text=True)

    sheet.append(HEADERS)
    for column, width in enumerate(WIDTHS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = top_wrap
        sheet.column_dimensions[get_column_letter(column)].width = width

    for index, row in enumerate(rows, start=2):
        verdict = row.get("verdict", "")
        label = row.get("ground_truth", "") or ""
        corrected = row.get("corrected", "") or ""
        note = (row.get("note", "") or "").strip()

        if verdict in NO_TRANSCRIPTION:
            label_value: Any = label
            corrected_value: Any = ""
            no_breaks = with_breaks = None
            reason = NO_TRANSCRIPTION[verdict]
            note = f"{reason} — {note}" if note else reason
        else:
            left, right = diff_spans(label, corrected)
            label_value = _rich(left, red)
            corrected_value = _rich(right, blue)
            no_breaks = accuracy(label, corrected, keep_breaks=False)
            with_breaks = accuracy(label, corrected, keep_breaks=True)

        sheet.cell(row=index, column=1, value=row.get("post_id", ""))

        image_cell = sheet.cell(row=index, column=2, value=row.get("image", ""))
        link = row.get("post_link", "")
        if link:
            image_cell.hyperlink = link
            image_cell.font = Font(color="FF0563C1", underline="single")

        sheet.cell(row=index, column=3, value=row.get("caption", ""))
        sheet.cell(row=index, column=4, value=label_value)
        sheet.cell(row=index, column=5, value=corrected_value)

        for column, value in ((6, no_breaks), (7, with_breaks)):
            cell = sheet.cell(row=index, column=column, value=value)
            # A real number with a percent format, never the string "78.57%":
            # a text column cannot be averaged, and averaging is the first
            # thing anyone does with this sheet.
            cell.number_format = "0.00%"
            cell.alignment = Alignment(vertical="top", horizontal="right")

        sheet.cell(row=index, column=8, value=note)

        for column in range(1, 9):
            cell = sheet.cell(row=index, column=column)
            cell.border = border
            if column not in (6, 7):
                cell.alignment = top_wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(1, len(rows)) + 1}"
    return book
