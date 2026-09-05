import io

import pytest

pytest.importorskip("openpyxl")

from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402

from app.core import evalsheet  # noqa: E402

# Their worked example, from the format screenshot: a four-line poem where the
# label is simplified and the corrected text traditional — six characters differ.
SIMPLIFIED = "花开花落又一秋\n时光匆匆似水流\n多少往事随风去\n悠悠岁月染白头"
TRADITIONAL = "花開花落又一秋\n時光匆匆似水流\n多少往事隨風去\n悠悠歲月染白頭"


def row(**overrides):
    base = {
        "post_id": "UzpfSTEyMzQ1Njc4OTA6Vks6OTg3NjU0MzIxMA==",
        "image": "UzpfSTEyMzQ1Njc4OTA6Vks6OTg3NjU0MzIxMA==_0.jpg",
        "caption": SIMPLIFIED,
        "ground_truth": SIMPLIFIED,
        "corrected": TRADITIONAL,
        "note": "Caption dùng chữ giản thể",
        "verdict": "minor",
        "post_link": "https://www.facebook.com/permalink.php?story_fbid=1&id=2",
    }
    base.update(overrides)
    return base


def sheet_of(rows):
    import openpyxl

    buffer = io.BytesIO()
    evalsheet.build(rows).save(buffer)
    buffer.seek(0)
    return openpyxl.load_workbook(buffer, rich_text=True).active


def coloured(value) -> str:
    """The characters this cell highlights as differing."""
    if not isinstance(value, CellRichText):
        return ""
    return "".join(b.text for b in value if isinstance(b, TextBlock))


class TestAccuracy:
    def test_reproduces_their_worked_example(self):
        """78.57% and 80.65% are the numbers on their own screenshot."""
        without = evalsheet.accuracy(SIMPLIFIED, TRADITIONAL, keep_breaks=False)
        with_breaks = evalsheet.accuracy(SIMPLIFIED, TRADITIONAL, keep_breaks=True)
        assert round(without * 100, 2) == 78.57
        assert round(with_breaks * 100, 2) == 80.65

    def test_line_breaks_lift_the_score_by_padding_the_denominator(self):
        """Counting separators that never differ can only flatter the result."""
        without = evalsheet.accuracy(SIMPLIFIED, TRADITIONAL, keep_breaks=False)
        with_breaks = evalsheet.accuracy(SIMPLIFIED, TRADITIONAL, keep_breaks=True)
        assert with_breaks > without

    def test_identical_text_is_a_hundred_percent(self):
        assert evalsheet.accuracy(SIMPLIFIED, SIMPLIFIED, keep_breaks=False) == 1.0

    def test_both_empty_is_a_hundred_percent(self):
        assert evalsheet.accuracy("", "", keep_breaks=False) == 1.0

    def test_one_empty_is_zero(self):
        assert evalsheet.accuracy("年歲漸長", "", keep_breaks=False) == 0.0

    def test_the_denominator_is_the_longer_side(self):
        """Their Task.xlsx convention, not CER — divide by max, not reference."""
        # 1 char vs 4: distance 3, max 4 -> 25%. CER would divide by 1 and clamp.
        assert evalsheet.accuracy("一", "一二三四", keep_breaks=False) == 0.25


class TestDiff:
    def test_only_the_differing_characters_are_coloured(self):
        sheet = sheet_of([row()])
        assert coloured(sheet.cell(row=2, column=4).value) == "开时随风岁头"
        assert coloured(sheet.cell(row=2, column=5).value) == "開時隨風歲頭"

    def test_an_insertion_does_not_paint_the_rest_of_the_line(self):
        """Position-by-position comparison would mark everything after it."""
        left, right = evalsheet.diff_spans("年歲漸長", "年歲漸增長")
        assert "".join(t for t, d in right if d) == "增"
        assert "".join(t for t, d in left if d) == ""

    def test_identical_text_is_left_as_a_plain_string(self):
        sheet = sheet_of([row(corrected=SIMPLIFIED, verdict="correct")])
        value = sheet.cell(row=2, column=4).value
        assert not isinstance(value, CellRichText)
        assert value == SIMPLIFIED

    def test_line_breaks_survive_the_diff(self):
        sheet = sheet_of([row()])
        assert "\n" in str(sheet.cell(row=2, column=4).value)


class TestSheet:
    def test_the_eight_columns_in_their_order(self):
        sheet = sheet_of([row()])
        assert [c.value for c in sheet[1]] == evalsheet.HEADERS

    def test_accuracies_are_numbers_not_percent_strings(self):
        sheet = sheet_of([row()])
        cell = sheet.cell(row=2, column=6)
        assert isinstance(cell.value, float)
        assert cell.number_format == "0.00%"
        assert round(cell.value * 100, 2) == 78.57

    def test_the_image_links_to_the_post(self):
        sheet = sheet_of([row()])
        cell = sheet.cell(row=2, column=2)
        assert cell.hyperlink.target.endswith("story_fbid=1&id=2")

    def test_a_missing_post_link_leaves_the_cell_plain(self):
        sheet = sheet_of([row(post_link="")])
        assert sheet.cell(row=2, column=2).hyperlink is None

    @pytest.mark.parametrize("verdict", ["unreadable", "not_an_image"])
    def test_unjudgeable_images_are_left_unscored(self, verdict):
        """0% would drag the average down with images nobody could judge."""
        sheet = sheet_of([row(verdict=verdict, corrected="")])
        assert sheet.cell(row=2, column=5).value in (None, "")
        assert sheet.cell(row=2, column=6).value is None
        assert sheet.cell(row=2, column=7).value is None

    def test_the_reason_moves_into_the_note(self):
        """This format has no verdict column, so it must not be lost."""
        sheet = sheet_of([row(verdict="unreadable", corrected="", note="mờ quá")])
        note = sheet.cell(row=2, column=8).value
        assert "Không đọc được" in note and "mờ quá" in note

    def test_a_reason_stands_alone_when_there_is_no_note(self):
        sheet = sheet_of([row(verdict="not_an_image", corrected="", note="")])
        assert sheet.cell(row=2, column=8).value == "Không phải ảnh thư pháp"

    def test_the_header_is_frozen_and_filterable(self):
        sheet = sheet_of([row(), row()])
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:H3"

    def test_an_empty_export_still_produces_a_usable_sheet(self):
        sheet = sheet_of([])
        assert [c.value for c in sheet[1]] == evalsheet.HEADERS
        assert sheet.max_row == 1
